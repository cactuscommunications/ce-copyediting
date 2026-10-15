from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
import re


def _sanitize_sheet_name(name: str) -> str:
    """
    Convert a filename into a valid Excel worksheet name.

    Excel sheet names:
    - Maximum 31 characters
    - Cannot contain: : \\ / ? * [ ]
    """
    name = re.sub(r'[:\\/?*\[\]]', "_", name)

    # Remove extension
    name = name.rsplit(".", 1)[0]

    # Excel does not allow empty sheet names
    if not name:
        name = "Sheet"

    return name[:31]


def _get_unique_sheet_name(wb, filename: str) -> str:
    """
    Generate a unique worksheet name from a filename.
    """
    base_name = _sanitize_sheet_name(filename)

    sheet_name = base_name
    counter = 1

    while sheet_name in wb.sheetnames:
        suffix = f"_{counter}"
        sheet_name = f"{base_name[:31 - len(suffix)]}{suffix}"
        counter += 1

    return sheet_name


def export_reference_edits_to_excel(
    references_by_file,
    output_file="reference_edits.xlsx"
):
    """
    Export reference edits into a single Excel workbook.

    Parameters
    ----------
    references_by_file : dict
        Mapping of filename -> list of reference dictionaries.

        Example:
        {
            "article_001.docx": [
                {
                    "original": original_string,
                    "edited": edited_string,
                    "edits": edits
                },
                ...
            ],
            "article_002.docx": [
                ...
            ]
        }

    output_file : str
        Output Excel file path.

    Each element in `edits` is expected to have:
        e.o_str
        e.c_str
    """

    wb = Workbook()

    # Remove the default worksheet
    default_ws = wb.active
    wb.remove(default_ws)

    # ----------------------------------------------------------
    # Create one worksheet per file
    # ----------------------------------------------------------

    for filename, references in references_by_file.items():

        sheet_name = _get_unique_sheet_name(wb, filename)
        ws = wb.create_sheet(title=sheet_name)

        # ------------------------------------------------------
        # Headers
        # ------------------------------------------------------

        headers = [
            "Original String",
            "Edited String",
            "Edit",
        ]

        ws.append(headers)

        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

        # ------------------------------------------------------
        # Data
        # ------------------------------------------------------
        if not references:
            continue
        for reference in references:

            original = reference["original"]
            edited = reference["edited"]
            edits = reference["edits"]

            errant_edits_list = [
                f'"{e.o_str}" to "{e.c_str}"'
                for i, e in enumerate(edits)
            ]

            # No edits
            if not errant_edits_list:
                ws.append([
                    original,
                    edited,
                    ""
                ])

            else:
                # First edit gets original/edited strings.
                # Remaining edits get blank values for those columns.
                for i, edit in enumerate(errant_edits_list):

                    ws.append([
                        original if i == 0 else "",
                        edited if i == 0 else "",
                        edit
                    ])

            # Empty row after each reference
            ws.append([])

        # ------------------------------------------------------
        # Formatting
        # ------------------------------------------------------

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True
                )

        ws.column_dimensions["A"].width = 50
        ws.column_dimensions["B"].width = 50
        ws.column_dimensions["C"].width = 90

        # Freeze header
        ws.freeze_panes = "A2"

        # Auto-filter
        ws.auto_filter.ref = ws.dimensions

    # ----------------------------------------------------------
    # Save workbook
    # ----------------------------------------------------------

    wb.save(output_file)